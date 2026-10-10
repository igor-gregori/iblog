"""Os leitores, testados com arquivos pequenos no mesmo formato dos originais (sem rede)."""

import json
import zipfile
from datetime import datetime

import openpyxl
import pandas as pd
import pytest

from comprar_ou_alugar.fontes import b3_ifix, bcb_sgs, cvm_fii, fipezap, ibge_sidra, tesouro


def aba_fipezap(planilha, nome, meses):
    """Cria uma aba no desenho da planilha do FipeZAP. `meses`: lista de (data, valor base)."""
    aba = planilha.create_sheet(nome)
    for coluna, titulo in fipezap.TITULOS_ESPERADOS.items():
        aba.cell(row=3, column=coluna + 1, value=titulo)
    aba.cell(row=4, column=2, value="Data")
    for linha, (data, base) in enumerate(meses, start=5):
        aba.cell(row=linha, column=2, value=data)
        for bloco, coluna in enumerate(fipezap.BLOCOS.values()):
            for quartos in range(5):
                valor = "." if base is None else base + 100 * bloco + quartos
                aba.cell(row=linha, column=coluna + quartos + 1, value=valor)
    return aba


@pytest.fixture
def planilha_fipezap(tmp_path):
    planilha = openpyxl.Workbook()
    resumo = planilha.active
    resumo.title = "Resumo"
    resumo.cell(row=8, column=2, value="Cidade A")
    resumo.cell(row=8, column=3, value="MG")
    planilha.create_sheet("Aux")
    meses = [(datetime(2020, 1, 1), None), (datetime(2020, 2, 1), 1000), (datetime(2020, 3, 1), 2000)]
    aba_fipezap(planilha, fipezap.NACIONAL, meses)
    aba_fipezap(planilha, "Cidade A", meses)
    caminho = tmp_path / "fipezap.xlsx"
    planilha.save(caminho)
    return caminho


def test_fipezap(planilha_fipezap):
    tabela = fipezap.ler(planilha_fipezap)

    assert set(tabela["cidade"]) == {"Cidade A", fipezap.NACIONAL}
    # Janeiro só tinha "." (sem dado) e some; sobram 2 meses × 5 recortes de dormitórios.
    cidade = tabela[tabela["cidade"] == "Cidade A"]
    assert len(cidade) == 10
    assert set(cidade["uf"]) == {"MG"}
    assert tabela[tabela["cidade"] == fipezap.NACIONAL]["uf"].isna().all()

    fevereiro = cidade[(cidade["data"] == "2020-02-01")].set_index("quartos")
    assert fevereiro.loc["total", "venda_indice"] == 1000
    assert fevereiro.loc["total", "venda_preco_m2"] == 1100
    assert fevereiro.loc["3", "aluguel_preco_m2"] == 1303
    assert fevereiro.loc["4", "rentabilidade_aluguel_mes"] == 1404


def test_fipezap_avisa_quando_o_formato_muda(planilha_fipezap):
    planilha = openpyxl.load_workbook(planilha_fipezap)
    planilha["Cidade A"].cell(row=3, column=18, value="Outra coisa")
    planilha.save(planilha_fipezap)

    with pytest.raises(ValueError, match="mudou de formato"):
        fipezap.ler(planilha_fipezap)


def test_b3_ifix(tmp_path):
    vazio = {f"rateValue{mes}": None for mes in range(1, 13)}
    resposta = {
        "results": [
            {"day": 2, **vazio, "rateValue1": "1.001,11", "rateValue3": "3.771,85"},
            {"day": 30, **vazio, "rateValue1": "", "rateValue12": "1.000,00"},
        ]
    }
    (tmp_path / "ifix_2011.json").write_text(json.dumps(resposta))

    tabela = b3_ifix.ler(tmp_path)

    assert tabela["data"].tolist() == [
        pd.Timestamp("2011-01-02"),
        pd.Timestamp("2011-03-02"),
        pd.Timestamp("2011-12-30"),
    ]
    assert tabela["ifix"].tolist() == [1001.11, 3771.85, 1000.0]


def test_ibge_sidra(tmp_path):
    def linha(variavel, mes, valor):
        return {"D2C": variavel, "D3C": mes, "V": valor}

    linhas = [
        {"D2C": "Variável (Código)", "D3C": "Mês (Código)", "V": "Valor"},
        linha("2266", "202607", "7657.73"),
        linha("63", "202607", "0.07"),
        linha("2266", "202608", "7633.23"),
        linha("63", "202608", "-0.32"),
        linha("2266", "202609", "..."),
    ]
    caminho = tmp_path / "ipca.json"
    caminho.write_text(json.dumps(linhas))

    tabela = ibge_sidra.ler(caminho)

    assert tabela["data"].tolist() == [pd.Timestamp("2026-07-01"), pd.Timestamp("2026-08-01")]
    assert tabela["ipca_indice"].tolist() == [7657.73, 7633.23]
    assert tabela["ipca_var_mes"].tolist() == pytest.approx([0.0007, -0.0032])


def test_bcb_sgs(tmp_path):
    envelope = """<?xml version="1.0" encoding="utf-8"?>
<soapenv:Envelope xmlns:soapenv="http://schemas.xmlsoap.org/soap/envelope/">
<soapenv:Body><ns1:getValoresSeriesXMLResponse xmlns:ns1="http://publico.ws.casosdeuso.sgs.pec.bcb.gov.br">
<getValoresSeriesXMLReturn>&lt;?xml version='1.0' encoding='ISO-8859-1'?&gt;
&lt;SERIES&gt;&lt;SERIE ID='1'&gt;
&lt;ITEM&gt;&lt;DATA&gt;1/2026&lt;/DATA&gt;&lt;VALOR&gt;{janeiro}&lt;/VALOR&gt;&lt;/ITEM&gt;
&lt;ITEM&gt;&lt;DATA&gt;2/2026&lt;/DATA&gt;&lt;VALOR&gt;{fevereiro}&lt;/VALOR&gt;&lt;/ITEM&gt;
&lt;ITEM&gt;&lt;DATA&gt;3/2026&lt;/DATA&gt;&lt;VALOR&gt;&lt;/VALOR&gt;&lt;/ITEM&gt;
&lt;/SERIE&gt;&lt;/SERIES&gt;</getValoresSeriesXMLReturn>
</ns1:getValoresSeriesXMLResponse></soapenv:Body></soapenv:Envelope>"""
    for nome, codigo in bcb_sgs.SERIES.items():
        valores = (
            dict(janeiro="786.62", fevereiro="790.08") if nome == "ivgr" else dict(janeiro="1.16", fevereiro="1.00")
        )
        (tmp_path / f"{codigo}_{nome}.xml").write_text(envelope.format(**valores))

    tabela = bcb_sgs.ler(tmp_path)

    assert tabela["data"].tolist() == [pd.Timestamp("2026-01-01"), pd.Timestamp("2026-02-01")]
    assert tabela["cdi_var_mes"].tolist() == pytest.approx([0.0116, 0.0100])  # percentual vira fração
    assert tabela["ivgr"].tolist() == [786.62, 790.08]  # índice fica como está


def test_bcb_sgs_data_diaria():
    assert bcb_sgs._data("30/9/2026") == pd.Timestamp("2026-09-30")


def test_cvm_fii(tmp_path):
    geral = (
        "CNPJ_Fundo;Data_Referencia;Versao;Nome_Fundo;Segmento_Atuacao;Mercado_Negociacao_Bolsa\n"
        "11.111.111/0001-11;2020-01-01;1;FUNDO UM;Shoppings;S\n"
        "11.111.111/0001-11;2020-01-01;2;FUNDO UM;Shoppings;S\n"
        "22.222.222/0001-22;2020-01-01;1;FUNDO DOIS;Outros;N\n"
    )
    complemento = (
        "CNPJ_Fundo;Data_Referencia;Versao;Total_Numero_Cotistas;Patrimonio_Liquido;Percentual_Dividend_Yield_Mes\n"
        "11.111.111/0001-11;2020-01-01;1;5000;1000000;0.99\n"
        "11.111.111/0001-11;2020-01-01;2;5000;1000000;0.006\n"
        "22.222.222/0001-22;2020-01-01;1;1;500000;0\n"
    )
    with zipfile.ZipFile(tmp_path / "inf_mensal_fii_2020.zip", "w") as pacote:
        pacote.writestr("inf_mensal_fii_geral_2020.csv", geral.encode("latin-1"))
        pacote.writestr("inf_mensal_fii_complemento_2020.csv", complemento.encode("latin-1"))

    tabela = cvm_fii.ler(tmp_path)

    assert list(tabela.columns) == list(cvm_fii.COLUNAS.values())
    assert len(tabela) == 2
    um = tabela.set_index("cnpj").loc["11.111.111/0001-11"]
    assert um["dividend_yield_mes"] == 0.006  # ficou a versão 2 do informe
    assert um["negociado_em_bolsa"]
    assert not tabela.set_index("cnpj").loc["22.222.222/0001-22", "negociado_em_bolsa"]


def test_tesouro(tmp_path):
    caminho = tmp_path / "tesouro.csv"
    caminho.write_text(
        "Tipo Titulo;Data Vencimento;Data Base;Taxa Compra Manha;Taxa Venda Manha;PU Compra Manha;PU Venda Manha;PU Base Manha\n"
        "Tesouro IPCA+;15/05/2035;02/01/2026;7,00;7,10;1;1;1\n"
        "Tesouro IPCA+;15/08/2040;02/01/2026;8,00;8,10;1;1;1\n"
        "Tesouro IPCA+;15/05/2035;05/01/2026;6,50;6,60;1;1;1\n"
        "Tesouro Selic;01/03/2031;02/01/2026;0,09;0,10;1;1;1\n"
    )

    tabela = tesouro.ler(caminho)

    # Dia 2: média de 7 e 8 = 7,5. Dia 5: 6,5. Mês: média dos dois dias = 7,0.
    assert tabela["data"].tolist() == [pd.Timestamp("2026-01-01")]
    assert tabela["tesouro_ipca_taxa_aa"].tolist() == pytest.approx([0.07])
