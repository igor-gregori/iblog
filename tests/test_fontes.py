"""Os leitores, testados com arquivos pequenos no mesmo formato dos originais (sem rede)."""

import json
import zipfile
from xml.sax.saxutils import escape

import openpyxl
import pandas as pd
import pytest

from juro_negativo.fontes import (
    afirmacoes,
    b3_ibovespa,
    banco_mundial,
    bcb_sgs,
    bis,
    bitcoin,
    fed,
    ibge_sidra,
    sp500,
    tesouro,
)


def resposta_do_sgs(itens: list[tuple[str, str]]) -> str:
    """Um envelope SOAP como o do Banco Central: o XML da série vem dentro, escapado como texto."""
    interno = "<?xml version='1.0' encoding='ISO-8859-1'?><SERIES><SERIE ID='1'>"
    interno += "".join(
        f"<ITEM><DATA>{data}</DATA><VALOR>{valor}</VALOR><BLOQUEADO>false</BLOQUEADO></ITEM>" for data, valor in itens
    )
    interno += "</SERIE></SERIES>"
    return (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<soapenv:Envelope xmlns:soapenv="http://schemas.xmlsoap.org/soap/envelope/"><soapenv:Body>'
        '<ns1:getValoresSeriesXMLResponse xmlns:ns1="http://publico.ws.casosdeuso.sgs.pec.bcb.gov.br">'
        f"<getValoresSeriesXMLReturn>{escape(interno)}</getValoresSeriesXMLReturn>"
        "</ns1:getValoresSeriesXMLResponse></soapenv:Body></soapenv:Envelope>"
    )


def test_bcb_sgs(tmp_path):
    mensal = [("1/2020", "0.50"), ("2/2020", ""), ("3/2020", "2000000")]
    diaria = [("30/1/2020", "4.50"), ("31/1/2020", "4.25"), ("28/2/2020", "4.00")]
    for nome, codigo in bcb_sgs.SERIES.items():
        itens = diaria if nome == "selic_meta_aa" else mensal
        (tmp_path / f"{codigo}_{nome}.xml").write_text(resposta_do_sgs(itens))

    tabela = bcb_sgs.ler(tmp_path).set_index("data")

    assert list(tabela.columns) == list(bcb_sgs.SERIES)
    assert tabela.loc["2020-01-01", "selic_var_mes"] == pytest.approx(0.005)  # percentual vira fração
    assert pd.isna(tabela.loc["2020-02-01", "selic_var_mes"])  # valor vazio não entra
    assert tabela.loc["2020-03-01", "m2"] == pytest.approx(2.0)  # R$ mil → R$ bilhões
    assert tabela.loc["2020-03-01", "pib_12m"] == pytest.approx(2000.0)  # R$ milhões → R$ bilhões
    assert tabela.loc["2020-03-01", "salario_minimo"] == 2_000_000  # em reais, sem conversão
    # A série diária entra com o último dia de cada mês.
    assert tabela.loc["2020-01-01", "selic_meta_aa"] == pytest.approx(0.0425)
    assert tabela.loc["2020-02-01", "selic_meta_aa"] == pytest.approx(0.04)


def test_b3_ibovespa_traz_os_pontos_antigos_para_a_escala_de_hoje(tmp_path):
    vazio = {f"rateValue{mes}": None for mes in range(1, 13)}
    antes = {
        "results": [{"day": 28, **vazio, "rateValue2": "88.287,30"}, {"day": 3, **vazio, "rateValue3": "8.978,22"}]
    }
    muito_antes = {"results": [{"day": 9, **vazio, "rateValue2": "90.466,84"}, {"day": 10, **vazio, "rateValue2": ""}]}
    hoje = {"results": [{"day": 31, **vazio, "rateValue8": "109.522,88"}]}
    (tmp_path / "ibov_1994.json").write_text(json.dumps(muito_antes))
    (tmp_path / "ibov_1997.json").write_text(json.dumps(antes))
    (tmp_path / "ibov_2022.json").write_text(json.dumps(hoje))

    tabela = b3_ibovespa.ler(tmp_path).set_index("data")["ibov"]

    assert list(tabela.index) == list(pd.to_datetime(["1994-02-09", "1997-02-28", "1997-03-03", "2022-08-31"]))
    assert tabela["1994-02-09"] == pytest.approx(904.6684)  # duas divisões por 10 pela frente
    assert tabela["1997-02-28"] == pytest.approx(8828.73)  # uma divisão por 10 pela frente
    assert tabela["1997-03-03"] == pytest.approx(8978.22)
    assert tabela["2022-08-31"] == pytest.approx(109522.88)


@pytest.fixture
def pink_sheet(tmp_path):
    planilha = openpyxl.Workbook()
    aba = planilha.active
    aba.title = "Monthly Prices"
    aba.cell(row=5, column=2, value="Crude oil, average")
    aba.cell(row=5, column=3, value="Gold")
    aba.cell(row=6, column=3, value="($/troy oz)")
    for linha, (mes, ouro) in enumerate([("2026M07", 4073), ("2026M08", 4411), ("2026M09", "…")], start=7):
        aba.cell(row=linha, column=1, value=mes)
        aba.cell(row=linha, column=2, value=70.5)
        aba.cell(row=linha, column=3, value=ouro)
    caminho = tmp_path / "pink.xlsx"
    planilha.save(caminho)
    return caminho


def test_banco_mundial_ouro(pink_sheet):
    tabela = banco_mundial.ler_ouro(pink_sheet)

    # Setembro só tinha "…" (sem dado) e some.
    assert tabela["data"].tolist() == list(pd.to_datetime(["2026-07-01", "2026-08-01"]))
    assert tabela["ouro_usd"].tolist() == [4073, 4411]


def test_banco_mundial_avisa_quando_a_pink_sheet_muda(pink_sheet):
    planilha = openpyxl.load_workbook(pink_sheet)
    planilha["Monthly Prices"].cell(row=6, column=3, value="($/kg)")
    planilha.save(pink_sheet)

    with pytest.raises(ValueError, match="mudou de formato"):
        banco_mundial.ler_ouro(pink_sheet)


def test_banco_mundial_indicadores(tmp_path):
    def registro(indicador, pais, ano, valor):
        return {"indicator": {"id": indicador}, "countryiso3code": pais, "date": str(ano), "value": valor}

    respostas = {
        "demografia": [registro("SP.DYN.TFRT.IN", "BRA", 2022, 1.629), registro("SP.DYN.TFRT.IN", "BRA", 2025, None)],
        "idades": [
            registro("SP.POP.0004.MA", "BRA", 2025, 6_500_000),
            registro("SP.POP.80UP.FE", "BRA", 2025, 3_000_000),
        ],
        "economia": [registro("NY.GDP.MKTP.CD", "BRA", 2010, 2.2e12)],
    }
    for nome, registros in respostas.items():
        (tmp_path / f"{nome}.json").write_text(json.dumps([{"page": 1}, registros]))

    tabela = banco_mundial.ler_indicadores(tmp_path).set_index(["indicador", "pais", "ano"])["valor"]

    assert len(tabela) == 4  # o ano sem valor fica de fora
    assert tabela["fecundidade", "BRA", 2022] == 1.629
    assert tabela["idade_0004_homens", "BRA", 2025] == 6_500_000
    assert tabela["idade_80UP_mulheres", "BRA", 2025] == 3_000_000
    assert tabela["pib_dolares", "BRA", 2010] == 2.2e12


def pacote_do_fed(caminho, pacote, series: dict[str, list[tuple[str, str, str]]]):
    """Um zip como os do Data Download Program. `series`: nome → lista de (data, valor, status)."""
    corpo = "".join(
        f'<kf:Series FREQ="129" SERIES_NAME="{nome}">'
        + "<frb:Annotations><common:Annotation><common:AnnotationText>x</common:AnnotationText></common:Annotation></frb:Annotations>"
        + "".join(
            f'<frb:Obs OBS_STATUS="{status}" OBS_VALUE="{valor}" TIME_PERIOD="{data}" />' for data, valor, status in obs
        )
        + "</kf:Series>"
        for nome, obs in series.items()
    )
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<message:MessageGroup xmlns:message="http://www.SDMX.org/resources/SDMXML/schemas/v1_0/message" '
        'xmlns:common="http://www.SDMX.org/resources/SDMXML/schemas/v1_0/common" '
        'xmlns:frb="http://www.federalreserve.gov/structure/compact/common" '
        f'xmlns:kf="http://www.federalreserve.gov/structure/compact/{pacote}_{pacote}">'
        f"<frb:DataSet>{corpo}</frb:DataSet></message:MessageGroup>"
    )
    with zipfile.ZipFile(caminho / f"{pacote}.zip", "w") as arquivo:
        arquivo.writestr(f"{pacote}_data.xml", xml)


def test_fed(tmp_path):
    dinheiro = [("2020-04-30", "4856.4", "A"), ("2020-05-31", "16312.5", "A")]
    pacote_do_fed(tmp_path, "H6", {"M1.M": dinheiro, "M2.M": dinheiro, "M1.WM": [("2020-05-04", "1.0", "A")]})
    juros = [("2020-04-30", "0.05", "A"), ("2020-05-31", "-9999", "ND")]
    pacote_do_fed(tmp_path, "H15", {"RIFSPFF_N.M": juros, "RIFLGFCY10_XII_N.M": [("2020-05-31", "-0.44", "A")]})

    tabela = fed.ler(tmp_path).set_index("data")

    assert tabela.loc["2020-05-01", "m1"] == 16312.5  # a data vai para o primeiro dia do mês
    assert tabela.loc["2020-04-01", "fed_funds_aa"] == pytest.approx(0.0005)
    assert pd.isna(tabela.loc["2020-05-01", "fed_funds_aa"])  # observação marcada como "sem dado"
    assert tabela.loc["2020-05-01", "tips_10_anos_aa"] == pytest.approx(-0.0044)


def test_fed_avisa_quando_falta_uma_serie(tmp_path):
    pacote_do_fed(tmp_path, "H6", {"M1.M": [("2020-04-30", "1.0", "A")]})
    pacote_do_fed(tmp_path, "H15", {})

    with pytest.raises(ValueError, match="não encontradas"):
        fed.ler(tmp_path)


def test_sp500(tmp_path):
    caminho = tmp_path / "data.csv"
    caminho.write_text(
        "Date,SP500,Dividend,Earnings\n"
        "2023-05-01,4146.17,68.71,170.0\n"
        "2023-06-01,4345.37,0.0,0.0\n"
        "2026-12-01,0.0,0.0,0.0\n"
    )

    tabela = sp500.ler(caminho)

    assert len(tabela) == 2  # a linha sem preço some
    assert tabela["sp500_dividendo_12m"].iloc[0] == 68.71
    assert pd.isna(tabela["sp500_dividendo_12m"].iloc[1])  # dividendo ainda não divulgado


def base_do_bis(caminho, base, linhas: list[dict]):
    with zipfile.ZipFile(caminho / f"{base}.zip", "w") as arquivo:
        arquivo.writestr(f"{base}_csv_flat.csv", pd.DataFrame(linhas).to_csv(index=False))


def test_bis(tmp_path):
    def linha(frequencia, pais, periodo, valor, unidade="368: Per cent per year"):
        return {
            bis.FREQUENCIA: frequencia,
            bis.PAIS: pais,
            bis.PERIODO: periodo,
            bis.VALOR: valor,
            bis.UNIDADE: unidade,
        }

    base_do_bis(
        tmp_path,
        bis.BASE_JUROS,
        [
            linha("M: Monthly", "BR: Brazil", "2026-08", 14.0),
            linha("D: Daily", "BR: Brazil", "2026-08-31", 14.0),
            linha("M: Monthly", "NZ: New Zealand", "2026-08", 3.0),  # país fora da lista
            linha("M: Monthly", "XM: Euro area", "2026-08", None),
        ],
    )
    base_do_bis(
        tmp_path,
        bis.BASE_PRECOS,
        [
            linha("M: Monthly", "BR: Brazil", "2026-08", 280.5, bis.INDICE),
            linha("M: Monthly", "BR: Brazil", "2026-08", 4.2, "771: Year-on-year changes, in per cent"),
            linha("A: Annual", "BR: Brazil", "2025", 270.0, bis.INDICE),
        ],
    )

    juros, precos = bis.ler_juros(tmp_path), bis.ler_precos(tmp_path)

    assert juros.to_dict("records") == [{"pais": "Brasil", "data": pd.Timestamp("2026-08-01"), "juro_aa": 0.14}]
    assert precos.to_dict("records") == [{"pais": "Brasil", "data": pd.Timestamp("2026-08-01"), "precos_indice": 280.5}]


def test_tesouro(tmp_path):
    caminho = tmp_path / "tesouro.csv"
    caminho.write_text(
        "Tipo Titulo;Data Vencimento;Data Base;Taxa Compra Manha;Taxa Venda Manha\n"
        "Tesouro IPCA+;15/05/2035;01/08/2026;8,00;8,10\n"
        "Tesouro IPCA+;15/05/2045;01/08/2026;7,00;7,10\n"
        "Tesouro IPCA+;15/05/2035;02/08/2026;9,50;9,60\n"
        "Tesouro Selic;01/03/2031;02/08/2026;0,09;0,10\n"
    )

    tabela = tesouro.ler(caminho)

    # Dia 1: média de 8% e 7%. Dia 2: 9,5%. Mês: média dos dois dias.
    assert tabela["data"].tolist() == [pd.Timestamp("2026-08-01")]
    assert tabela["tesouro_ipca_taxa_aa"].iloc[0] == pytest.approx(0.085)


def test_bitcoin(tmp_path):
    caminho = tmp_path / "btc.json"
    pontos = [{"x": 1230940800, "y": 0.0}, {"x": 1282089600, "y": 0.07}, {"x": 1791072000, "y": 84743.72}]
    caminho.write_text(json.dumps({"status": "ok", "values": pontos}))

    tabela = bitcoin.ler(caminho)

    # Os dias sem mercado (preço zero) ficam de fora.
    assert tabela["data"].tolist() == list(pd.to_datetime(["2010-08-18", "2026-10-04"]))
    assert tabela["btc_usd"].tolist() == [0.07, 84743.72]


def test_ibge_sidra(tmp_path):
    caminho = tmp_path / "nascimentos.json"
    linhas = [{"D3C": "Ano (Código)", "V": "Valor"}, {"D3C": "2024", "V": "2421079"}, {"D3C": "2003", "V": "3408540"}]
    caminho.write_text(json.dumps([*linhas, {"D3C": "2025", "V": "..."}]))

    tabela = ibge_sidra.ler(caminho)

    assert tabela.to_dict("list") == {"ano": [2003, 2024], "nascimentos": [3408540, 2421079]}


def test_afirmacoes_viram_variacao_em_fracao():
    tabela = afirmacoes.ler().set_index(["pais", "inicio", "fim", "serie"])["variacao_citada"]

    assert len(tabela) == 36
    assert tabela.index.is_unique
    # "26,48 vezes" é uma alta de 2.548%; "483,59%" é 4,8359; "-1,4%" é -0,014.
    assert tabela["Brasil", "1995-01-01", "2022-09-01", "ouro"] == pytest.approx(25.48)
    assert tabela["Brasil", "1998-01-01", "2008-01-01", "selic"] == pytest.approx(4.8359)
    assert tabela["Brasil", "2020-01-01", "2022-09-01", "ibovespa"] == pytest.approx(-0.014)
